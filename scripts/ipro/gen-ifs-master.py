# /home/rokunet/projects/ilerihub-ipro/scripts/ipro/gen-ifs-master.py
# çalıştırma:  python3 scripts/ipro/gen-ifs-master.py ~/ipro-mas-export/TEZGAH.xlsx
# gereksinim:  pip install openpyxl --break-system-packages

import sys, openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

SRC = sys.argv[1] if len(sys.argv) > 1 else "TEZGAH.xlsx"
OUT = "IPRO_IFS_Master.xlsx"

# ---------------------------------------------------------------- iş merkezleri
WC = [
 ("WCN01","CNC TORNA","WCN","CNC / Talasli Imalat"),
 ("WCN02","CNC ISLEME MERKEZI","WCN","CNC / Talasli Imalat"),
 ("WCN03","REVOLVER TORNA","WCN","CNC / Talasli Imalat"),
 ("WCN04","MATKAP","WCN","CNC / Talasli Imalat"),
 ("WCN05","DIS ACMA / OVALAMA","WCN","CNC / Talasli Imalat"),
 ("WCN06","BOY KESME / TESTERE","WCN","CNC / Talasli Imalat"),
 ("WCN07","TASLAMA","WCN","CNC / Talasli Imalat"),
 ("WDT01","DAIRE TESTERE","WDT","Daire Testere / Boru Bukum"),
 ("WDT02","CNC BORU BUKUM","WDT","Daire Testere / Boru Bukum"),
 ("WDT03","BORU UC ISLEME","WDT","Daire Testere / Boru Bukum"),
 ("WDT04","LAMA / TEL SEKILLENDIRME","WDT","Daire Testere / Boru Bukum"),
 ("WDT05","CEMBERLEME","WDT","Daire Testere / Boru Bukum"),
 ("WDT06","TASLAMA","WDT","Daire Testere / Boru Bukum"),
 ("WPH01","EKSANTRIK PRES","WPH","Preshane"),
 ("WPH02","HIDROLIK PRES","WPH","Preshane"),
 ("WPH03","CNC ABKANT BUKUM","WPH","Preshane"),
 ("WPH04","GIYOTIN MAKAS","WPH","Preshane"),
 ("WLZ01","LAZER KESIM","WLZ","Lazer"),
 ("WKR01","KAYNAK ROBOTU 1 (GEDIK OTC-1)","WKR","Kaynak Robot"),
 ("WKR02","KAYNAK ROBOTU 2 (GEDIK OTC-2)","WKR","Kaynak Robot"),
 ("WKR03","KAYNAK ROBOTU 3 (FANUC-1)","WKR","Kaynak Robot"),
 ("WKR04","KAYNAK ROBOTU 4 (FANUC-2)","WKR","Kaynak Robot"),
 ("WKR05","8014 ROBOT HATTI","WKR","Kaynak Robot"),
 ("WKR06","NACHI-GEKA","WKR","Kaynak Robot"),
 ("WKR07","PRES (ROBOT ALANI)","WKR","Kaynak Robot"),
 ("WKY01","GAZALTI KAYNAK","WKY","Manuel Kaynak"),
 ("WKY02","YARI OTOMATIK ROBOT KAYNAK","WKY","Manuel Kaynak"),
 ("WKY03","ZIMPARA / TASLAMA","WKY","Manuel Kaynak"),
 ("WKY04","KUMLAMA","WKY","Manuel Kaynak"),
 ("WKY05","KURT AGZI ACMA","WKY","Manuel Kaynak"),
 ("WKY06","MARKALAMA PRESI","WKY","Manuel Kaynak"),
 ("WPN01","PUNTA KAYNAK","WPN","Punta Kaynak"),
 ("WPN02","INDUKSIYON","WPN","Punta Kaynak"),
 ("WSK01","SASI GAZALTI KAYNAK","WSK","Sasi Kaynak"),
 ("WSK02","SASI PUNTA","WSK","Sasi Kaynak"),
 ("WSK03","SASI ISLEME","WSK","Sasi Kaynak"),
 ("WSK04","TOZ TOPLAMA","WSK","Sasi Kaynak"),
 ("WKP01","FREZE","WKP","Kaliphane"),
 ("WKP02","CNC ISLEME MERKEZI","WKP","Kaliphane"),
 ("WKP03","EREZYON","WKP","Kaliphane"),
 ("WKP04","MATKAP","WKP","Kaliphane"),
 ("WKP05","TORNA","WKP","Kaliphane"),
 ("WKP06","TESTERE","WKP","Kaliphane"),
 ("WKP07","TAKIM BILEME","WKP","Kaliphane"),
 ("WPE01","ENJEKSIYON 400T","WPE","Plastik Enjeksiyon"),
 ("WPE02","ENJEKSIYON 200T","WPE","Plastik Enjeksiyon"),
 ("WPE03","ENJEKSIYON 150T","WPE","Plastik Enjeksiyon"),
 ("WPE04","ENJEKSIYON 120T","WPE","Plastik Enjeksiyon"),
 ("WPE05","ENJEKSIYON 85T","WPE","Plastik Enjeksiyon"),
 ("WPE06","ENJEKSIYON 50T","WPE","Plastik Enjeksiyon"),
 ("WPE07","BOYAMA","WPE","Plastik Enjeksiyon"),
 ("WPE08","TAMPON BASKI","WPE","Plastik Enjeksiyon"),
 ("WPE09","ENJ. MONTAJ","WPE","Plastik Enjeksiyon"),
 ("WPE10","ENJ. PAKETLEME","WPE","Plastik Enjeksiyon"),
 ("WMM01","EL FREN MONTAJ HATTI","WMM","Mekanik Montaj"),
 ("WMM02","MONTAJ HATTI","WMM","Mekanik Montaj"),
 ("WMM03","MONTAJ TEZGAHI","WMM","Mekanik Montaj"),
 ("WMM04","EL GAZI KITLEK","WMM","Mekanik Montaj"),
 ("WMM05","MEK. MONTAJ PAKETLEME","WMM","Mekanik Montaj"),
 ("WKM01","KILIT MONTAJ HATTI","WKM","Kilit Montaj"),
 ("WKM02","KILIT PAKETLEME","WKM","Kilit Montaj"),
 ("WKM03","PSA MONTAJ HATTI","WKM","Kilit Montaj"),
 ("WHM01","HAPPICH MONTAJ HATTI","WHM","Happich Montaj"),
 ("WHM02","HAPPICH TEZGAH","WHM","Happich Montaj"),
 ("WHM03","HAPPICH PAKETLEME","WHM","Happich Montaj"),
 ("WPK01","GENEL PAKETLEME","WPK","Paketleme"),
 ("WPK02","DIREKSIYON KOLONU MONTAJ","WPK","Paketleme"),
 ("WPK03","DIREKSIYON KOLONU PAKETLEME","WPK","Paketleme"),
 ("WPK04","2197 / 2489 HATTI","WPK","Paketleme"),
 ("WPK05","MARKALAMA","WPK","Paketleme"),
]

# ------------------------------------------------- tezgah kodu -> is merkezi
M = {}
def put(wc, *kodlar):
    for k in kodlar: M[k] = wc

put("WCN01","CN01","CN02","CN03")
put("WCN02","CN17","CN19","CN21")
put("WCN03","CN16")
put("WCN04","CN08","CN09","CN10","CN11","CN12","CN13","CN14")
put("WCN05","CN18")
put("WCN06","CN04","CN20")
put("WCN07","CN15")
put("WDT01","DT02","DT08","DT11")
put("WDT02","DT03","DT06","DT12")
put("WDT03","DT04","DT07")
put("WDT04","DT05","DT10")
put("WDT05","DT14")
put("WDT06","DT15")
put("WPH01","PH04","PH06","PH07","PH08","PH09","PH10","PH11","PH12","PH15")
put("WPH02","PH01","PH02","PH03","PH05","PH20")
put("WPH03","PH19","PH21")
put("WPH04","PH13")
put("WLZ01","LZ01","LZ02")
put("WKR01","KR01-1","KR01-2","KR01-3","KR01-4","KR01-5","KR01-6")
put("WKR02","KR02-1","KR02-2","KR02-3","KR02-4","KR02-5","KR02-6")
put("WKR03","KR03-1","KR03-2","KR03-3","KR03-4","KR03-5","KR03-6")
put("WKR04","KR04-1","KR04-2","KR04-3","KR04-4","KR04-5","KR04-6")
put("WKR05","KR09","0158")
put("WKR06","KR10")
put("WKR07","KR05")
put("WKY01","KH01","KH02","KH03","KH04","KH05","KH07","KH16","KH17")
put("WKY02","KH12","KH13")
put("WKY03","KH08","KH09","KH14","KH15")
put("WKY04","KH10")
put("WKY05","KH06")
put("WKY06","KH11")
put("WPN01","KH26","KH27","KH28","KH31","KH32","KH37")
put("WPN02","KH29")
put("WSK01","SK04","SK05","SK06")
put("WSK02","SK01","SK03")
put("WSK03","SK02","SK07","SK08")
put("WSK04","SK09")
put("WKP01","KP01","KP02","KP03","KP07")
put("WKP02","KP09")
put("WKP03","KP10","KP11")
put("WKP04","KP04","KP05","KP08","KP12")
put("WKP05","KP15","KP16")
put("WKP06","KP06","KP13","KP14")
put("WKP07","KP17")
put("WPE01","PE01"); put("WPE02","PE02"); put("WPE03","PE03")
put("WPE04","PE04"); put("WPE05","PE05"); put("WPE06","PE06")
put("WPE07","PE21","PE25")
put("WPE08","PE22")
put("WPE09","PE23","PE24","PE26")
put("WPE10","PE08")
put("WMM01","MM202","MM230","MM233")
put("WMM02","MM30","MM150","MM151")
put("WMM03","MM01","MM154","0156","0157")
put("WMM04","MM210")
put("WMM05","MM03","MM152","MM153","MM155","MM205","MM211")
put("WKM01","KM01","KM02","KM03","KM04","KM05","KM06")
put("WKM02","KM07","KM08","KM09","KM010","KM011","KM012")
put("WKM03","MM59","MM60","MM61","MM62","MM63","MM64",
            "MM65","MM66","MM67","MM68","MM69","MM70")
put("WHM01","HM01","HM02")
put("WHM02","HM15","HM16")
put("WHM03","HM17","HM18")
put("WPK01","PK01","PK02","PK03","PK04")
put("WPK02","PK06","PK12")
put("WPK03","PK07","PK15")
put("WPK04","PK08","PK13","PK14")
put("WPK05","PK05","KM21","MM26","KM013","KM13")

# ---------------------------------- IPRO kodu -> IFS ResourceId (override)
IFS_ID = {}
for i in range(1, 7):
    IFS_ID[f"KR01-{i}"] = "KR01"
    IFS_ID[f"KR02-{i}"] = "KR02"
    IFS_ID[f"KR03-{i}"] = "KR03"
    IFS_ID[f"KR04-{i}"] = "KR04"
IFS_ID["0156"] = "MM71"
IFS_ID["0157"] = "MM72"
IFS_ID["0158"] = "KR11"
IFS_ID["MM26"]  = "PK05"
IFS_ID["KM013"] = "PK05"
IFS_ID["KM13"]  = "PK05"

NOT = {
 **{f"KR0{r}-{i}": "Robot ayni anda tek kapida calisir; 6 kapi IFS'te tek resource"
    for r in range(1, 5) for i in range(1, 7)},
 "0156": "Numerik MAS kodu; IFS'te anlamsiz",
 "0157": "Numerik MAS kodu; IFS'te anlamsiz",
 "0158": "Numerik MAS kodu; IFS'te anlamsiz",
 "MM26":  "Mukerrer markalama kaydi -> PK05 (mesai saatinde teyit)",
 "KM013": "Mukerrer markalama kaydi -> PK05 (mesai saatinde teyit)",
 "KM13":  "Mukerrer markalama kaydi -> PK05 (mesai saatinde teyit)",
}

# ------------------------------------------------------------------- oku
src = openpyxl.load_workbook(SRC, data_only=True)["Tezgah Listesi"]
aktif, fazla = [], []
for r in src.iter_rows(min_row=5, values_only=True):
    kat, durum, kod, ad = r[0], r[1], r[2], r[3]
    if not durum or not kod:
        continue
    (fazla if durum.startswith("FAZLALIK") else aktif).append((kat, durum, kod, ad))

eksik = [k for _, _, k, _ in aktif if k not in M]
if eksik:
    raise SystemExit(f"HATA - is merkezi atanmamis {len(eksik)} kod: {eksik}")

# ------------------------------------------------------------------- yaz
wb = Workbook(); wb.remove(wb.active)
H = Font(bold=True, color="FFFFFF")
F = PatternFill("solid", fgColor="1B4F72")

def sheet(name, headers, rows, widths):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for c in ws[1]: c.font, c.fill = H, F
    for row in rows: ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    return ws

sheet("IS_MERKEZI",
      ["WorkCenterNo","Description","Site","WorkCenterCode","LaborClass",
       "Department","Calendar","Utilization","SchedCapacity","UsageCode"],
      [[w, d, "ILER2", "Internal", lc, dept, "MAINT", 100, "Infinite", "Active"]
       for w, d, lc, dept in WC],
      [12, 34, 8, 15, 11, 28, 10, 12, 14, 11])

res, seen = [], set()
for kat, _, kod, ad in aktif:
    rid = IFS_ID.get(kod, kod)
    if rid in seen:
        continue
    seen.add(rid)
    res.append([rid, (ad or "").strip()[:35], M[kod], "ILER2", 100, "2026-08-01", kod])
res.sort(key=lambda x: (x[2], x[0]))
sheet("RESOURCE",
      ["ResourceId","Description","WorkCenterNo","Site","Efficiency",
       "StartDate","Kaynak MAS Kodu"],
      res, [12, 40, 14, 8, 11, 12, 16])

esl = [[kod, (ad or "").strip(), kat, M[kod], IFS_ID.get(kod, kod),
        "EVET" if kod in IFS_ID else "", NOT.get(kod, "")]
       for kat, _, kod, ad in aktif]
esl.sort(key=lambda x: (x[3], x[0]))
sheet("IPRO_ESLEME",
      ["IproTezgah.kod","TezgahAdi","MAS Grubu","WorkCenterNo",
       "ifsResourceId","Override?","Not"],
      esl, [15, 42, 26, 14, 14, 10, 52])

sheet("KAPSAM_DISI",
      ["TezgahKodu","TezgahAdi","MAS Grubu","Durum"],
      [[kod, (ad or "").strip(), kat, durum] for kat, durum, kod, ad in fazla],
      [12, 44, 30, 40])

wb.save(OUT)
print(f"{OUT} olustu")
print(f"  IS_MERKEZI   : {len(WC)}")
print(f"  RESOURCE     : {len(res)}")
print(f"  IPRO_ESLEME  : {len(esl)}")
print(f"  KAPSAM_DISI  : {len(fazla)}")
